"""Candidate closure, own event evidence and gait relations are separate."""
from dataclasses import replace
import pytest
from hardware.sensor_frame import DeviceLayout
from hardware.walking_preflight import PreparedDevice
from engine.modes.overground_running_processor import OvergroundRunningProcessor
from tests.test_overground_running import frame,processor


def capture(*,pulse=(),width=1,repeat=False,skip=(),bad=(),end=1200):
 p=processor(3,starting_foot='Left',min_contact_time=60)
 p=OvergroundRunningProcessor(p.config,PreparedDevice(p.device.layout,'running',-1,0,1000,bad_indices=bad))
 for n in range(end):
  if n in skip:continue
  bounds=[]
  for a,b,x in [(100,300,.3),(450,650,1.0),(800,1000,1.7)]:
   if a<=n<b:bounds.append((x-.08,x+.08))
  f=frame(p.device.layout,n,bounds)
  if n in pulse:
   bits=bytearray(f.contact_bits);bits[220:220+width]=b'\x01'*width;f=replace(f,contact_bits=bytes(bits))
  p.process(f,masked_contact=bool(bad) and n in pulse)
 return p


@pytest.mark.parametrize('pulse',[(250,),(350,),(1100,)])
def test_closed_isolated_nonstep_does_not_change_step_identity_support_or_duration(pulse):
 base=capture().summary();p=capture(pulse=pulse);s=p.summary()
 assert s['valid_steps']==base['valid_steps']==2 and s['valid_cycles']==base['valid_cycles']==1
 assert [c.side for c in p.contacts if c.confirmed]==['left','right','left']
 assert s['duration_s']==base['duration_s']
 assert [x['flight_s'] for x in s['steps']]==[x['flight_s'] for x in base['steps']]
 decision=next(x for x in s['candidate_decisions'] if x['first_sample']==pulse[0])
 assert decision['outcome']=='excluded_nonstep' and decision['closed_sample']==pulse[0]+3


@pytest.mark.parametrize('restriction',['repeated','wide','gap','bad','stop'])
def test_no_exclusion_without_single_sample_reliable_release(restriction):
 p=capture(pulse=(250,251) if restriction=='repeated' else(250,),width=2 if restriction=='wide' else 1,
           skip=(251,) if restriction=='gap' else(),bad=(219,) if restriction=='bad' else(),end=252 if restriction=='stop' else 1200)
 decisions=p.summary()['candidate_decisions']
 assert not any(x['outcome']=='excluded_nonstep' for x in decisions)
 assert p.summary()['valid_cycles']==0


def test_pending_return_preserves_original_onset_and_real_short_clear():
 p=processor(3,starting_foot='Left',min_contact_time=60)
 for n in range(700):
  bounds=[(.22,.38)] if 100<=n<400 else[]
  if n==250:bounds +=[(.9,.92)]
  if 252<=n<550:bounds +=[(.88,1.05)]
  p.process(frame(p.device.layout,n,bounds))
 assert len(p.contacts)==2
 assert p.contacts[1].start==250 and p.contacts[1].problem=='uncertain_short_clear'
 assert p.summary()['contacts'][1]['first_seen_sample']==250
 assert p.summary()['contacts'][1]['touch_sample'] is None
 assert not p.summary()['steps'][0]['flight_s']['valid']


def test_confirmation_does_not_make_unconfirmed_candidate_an_ab_identity():
 p=processor(3,starting_foot='Left',min_contact_time=60)
 for n in range(121):p.process(frame(p.device.layout,n,[(.2,.38)] if n>=100 else[]))
 assert p.contacts[0].label=='' and p.contacts[0].side=='unknown'
 assert p.summary()['contacts'][0]['confirmed_sample'] is None
 for n in range(121,170):p.process(frame(p.device.layout,n,[(.2,.38)]))
 assert p.contacts[0].label=='A' and p.contacts[0].side=='left'
 row=p.summary()['contacts'][0]
 assert row['touch_sample']==100 and row['confirmed_sample']==159


def test_cycle_contact_duration_does_not_depend_on_later_bad_contact():
 p=processor(8,starting_foot='Left',min_contact_time=60)
 p=OvergroundRunningProcessor(p.config,PreparedDevice(p.device.layout,'running',-1,0,1000,bad_indices=(576,)))
 for n in range(1100):
  bounds=[(.2,.4)] if 100<=n<300 else[(1.0,1.2)] if 450<=n<650 else[(6.0104,6.2)] if 800<=n<1000 else[]
  p.process(frame(p.device.layout,n,bounds),masked_contact=800<=n<1000)
 s=p.summary();assert s['cycles'][0]['contact_s']['valid']
 assert s['cycles'][0]['contact_s']['value']==pytest.approx(.2)
 assert not s['cycles'][0]['duration_s']['valid']


@pytest.mark.parametrize('pulse_sample', [1100, 1418])
def test_closed_nonstep_tail_restores_original_exit_deadline_but_waits_for_closure(pulse_sample):
 p=processor(1)
 for n in range(1700):
  bounds=[]
  if 100<=n<250:bounds=[(.15,.3)]
  if 400<=n<550:bounds=[(.45,.6)]
  if 700<=n<920:bounds=[(.82,.99 if 820<=n<840 else .96)]
  if n==pulse_sample:bounds +=[(p.positions[40],p.positions[40])]
  p.process(frame(p.device.layout,n,bounds))
  if p.finished_reason:break
 assert p.finished_reason=='passage_complete'
 assert n==max(1419,pulse_sample+3)
 assert p.last_occupied==919
 assert p.summary()['duration_s']==pytest.approx(.82)


def test_unresolved_tail_must_not_restore_old_terminal_evidence():
 p=processor(1)
 for n in range(1800):
  bounds=[]
  if 100<=n<250:bounds=[(.15,.3)]
  if 400<=n<550:bounds=[(.45,.6)]
  if 700<=n<920:bounds=[(.82,.99 if 820<=n<840 else .96)]
  if 1100<=n<1102:bounds +=[(p.positions[40],p.positions[40])]
  p.process(frame(p.device.layout,n,bounds))
 assert p.finished_reason is None and not p._exit_evidence
 assert any(x['outcome']=='unresolved' for x in p.summary()['candidate_decisions'])


def test_candidate_evidence_roundtrips_history_and_running_export(qtbot):
 from io import BytesIO
 from openpyxl import Workbook,load_workbook
 from data.subject_store import _report_detail,_report_from_detail
 from reporting.overground_running import detail_html,export_sheets
 p=capture(pulse=(250,350))
 report=_report_from_detail(_report_detail(p.build_report('manual')))
 assert report.running_summary==p.summary()
 assert '已排除非步候选' in detail_html(report)
 book=Workbook();export_sheets(book,report)
 stream=BytesIO();book.save(stream);book.close();stream.seek(0)
 saved=load_workbook(stream);rows=list(saved['Running candidate_decisions'].values);cols=rows[0]
 assert 'source' in cols and 'first_sample' in cols and 'contact_id' in cols
 assert {r[cols.index('first_sample')] for r in rows[1:]}=={250,350}
 assert all(r[cols.index('outcome')]=='excluded_nonstep' for r in rows[1:])
 saved.close()
