"""Deterministic 1000 Hz trajectories with independent geometry/time expectations."""
from dataclasses import replace
import time

import pytest
from config.walking_config import WalkingConfig
from config.test_config import config_from_dict
from config.config_validation import validate_runtime_config
from hardware.sensor_frame import DeviceLayout, SensorFrame, AcquisitionIssue
from hardware.walking_preflight import WalkingPreflight, PreparedDevice
from engine.modes.walking_processor import WalkingProcessor
from engine.walking_session import WalkingSession


def frame(layout, index, ranges=(), *, received=None, stream="test"):
    bits = bytes(int(any(low <= x <= high for low, high in ranges)) for x in layout.positions_m)
    return SensorFrame(stream, layout, index, index, received or time.perf_counter_ns(),
                       bits, bytes([1]) * layout.bit_count, b"")


def processor(n=3, **kwargs):
    layout = DeviceLayout.linear(n)
    return WalkingProcessor(WalkingConfig(**kwargs), PreparedDevice(layout, "test", -1, 0, 1000))


def walk(p, contacts, end_ms, *, skip=(), host_batch=False):
    for i in range(end_ms):
        if i in skip:
            continue
        ranges = [(position - .10, position + .10) for start, end, position in contacts if start <= i < end]
        p.process(frame(p.device.layout, i, ranges, received=(i // 16 + 1) * 16_000_000 if host_batch else i + 1))


CONTACTS = [(100, 800, .35), (600, 1300, .95), (1100, 1800, 1.55), (1600, 2300, 2.15)]


@pytest.mark.parametrize("n", [1, 3, 8, 12])
def test_preflight_requires_complete_clear_continuous_window(n):
    gate = WalkingPreflight()
    layout = DeviceLayout.linear(n)
    for i in range(999):
        gate.feed(frame(layout, i))
    assert not gate.ready(time.perf_counter_ns())
    gate.feed(frame(layout, 999))
    assert gate.ready(time.perf_counter_ns())
    assert gate.context.snapshot()["nominal_length_m"] == n
    gate.feed(frame(layout, 1000, [(.2, .3)]))
    assert not gate.ready(time.perf_counter_ns())
    for i in range(1001, 2001):
        gate.feed(frame(layout, i))
    assert gate.ready(time.perf_counter_ns())
    assert not gate.ready(gate.context.checked_monotonic_ns + gate.STALE_NS + 1)


def test_preflight_gap_mask_layout_and_stream_invalidate():
    gate = WalkingPreflight()
    layout = DeviceLayout.linear(3)
    for i in range(1000):
        gate.feed(frame(layout, i))
    gate.feed(frame(layout, 1001))
    assert not gate.context
    gate.feed(replace(frame(layout, 1002), valid_bits=bytes(layout.bit_count)))
    assert gate.healthy_samples == 0
    for i in range(1003, 2003):
        gate.feed(frame(layout, i))
    gate.feed(frame(DeviceLayout.linear(8), 2003))
    assert not gate.context
    gate.feed(frame(DeviceLayout.linear(8), 2004, stream="new"))
    assert gate.healthy_samples == 1


@pytest.mark.parametrize("n", [3, 8, 12])
def test_metric_geometry_and_device_time_are_length_independent(n):
    p = processor(n, stop_type="Software command", starting_foot="Left")
    walk(p, CONTACTS, 2400, host_batch=True)
    s = p.summary()
    assert p.origin == .1  # confirmation does not shift touch timestamp
    assert s["duration_s"] == pytest.approx(2.2)
    assert s["valid_steps"] == 3
    assert s["valid_cycles"] == 2
    assert s["mean_step_m"] == pytest.approx(.6, abs=.006)
    assert s["mean_stride_m"] == pytest.approx(1.2, abs=.006)
    assert s["mean_contact_s"] == pytest.approx(.7)
    assert s["cadence_per_min"] == pytest.approx(120)
    assert s["double_support_s"] == pytest.approx(.3)  # entry cycle .2 s, next cycle .4 s
    assert s["single_support_s"] == pytest.approx(.7)
    assert [c["side"] for c in s["contacts"]] == ["left", "right", "left", "right"]
    assert len(p.timeline[-1]["contact_bits"]) == n * 96


def test_seam_coordinates_are_not_global_index_times_pitch():
    p = processor(stop_type="Software command")
    walk(p, [(100, 400, .98), (600, 900, 1.58)], 1000)
    s = p.summary()
    assert len(s["contacts"]) == 2  # first foot straddles module seam
    first_bits = [x for x in p.positions if .88 <= x <= 1.08]
    second_bits = [x for x in p.positions if 1.48 <= x <= 1.68]
    expected = (second_bits[0] + second_bits[-1] - first_bits[0] - first_bits[-1]) / 2
    assert s["mean_step_m"] == pytest.approx(expected)
    assert all(c["side"] == "unknown" for c in s["contacts"])


def test_forward_reverse_auto_exit_and_manual_option():
    contacts = CONTACTS + [(2100, 2800, 2.75)]
    for reverse in (False, True):
        p = processor()
        route = [(a, b, 2.988 - x if reverse else x) for a, b, x in contacts]
        walk(p, route, 3400)
        assert p.finished_reason == "passage_complete"
        assert p.summary()["duration_s"] == pytest.approx(2.7)
    manual = processor(stop_type="Software command")
    walk(manual, contacts, 3400)
    assert manual.finished_reason is None
    middle = processor()
    walk(middle, CONTACTS[:2], 2200)
    assert middle.finished_reason is None  # clear in the middle is not exit


def test_gap_never_joins_steps_or_cycles_and_recovery_can_measure_new_contacts():
    p = processor(8, stop_type="Software command", starting_foot="Left")
    contacts = CONTACTS + [(2100, 2800, 2.75), (2600, 3300, 3.35), (3100, 3800, 3.95)]
    walk(p, contacts, 3900, skip=range(1400, 1410))
    s = p.summary()
    assert any(i["code"] == "frame_gap" for i in s["issues"])
    by_id = {c["id"]: c for c in s["contacts"]}
    for step in s["steps"]:
        assert by_id[step["from_id"]]["epoch"] == by_id[step["to_id"]]["epoch"]
    assert s["valid_steps"] > 0
    assert s["valid_cycles"] > 0
    assert s["passage_speed_m_s"] is None
    assert all(c["side"] == "unknown" for c in s["contacts"] if c["epoch"] > 0)


def test_stop_is_retained_in_passage_but_excluded_from_walking_metrics():
    p = processor(8, stop_type="Software command")
    contacts = [(100, 2800, .35), (2600, 3300, .95), (3100, 3800, 1.55), (3600, 4300, 2.15)]
    walk(p, contacts, 4500)
    s = p.summary()
    assert s["duration_s"] == pytest.approx(4.2)
    assert s["stops"]
    assert s["walking_speed_m_s"] > s["passage_speed_m_s"]
    assert s["cadence_per_min"] == pytest.approx(120)


def test_boundary_short_lane_and_ambiguous_contacts_do_not_fabricate_metrics():
    p = processor(1, stop_type="Software command")
    walk(p, [(100, 500, .02), (600, 900, .98)], 1100)
    s = p.summary()
    assert s["mean_step_m"] is None
    assert s["mean_stride_m"] is None
    assert s["excluded_contacts"] == 2
    q = processor()
    walk(q, [(100, 800, .3), (100, 800, .45)], 1000)
    assert q.summary()["mean_stride_m"] is None


def test_device_change_and_clock_reset_end_passage():
    p = processor()
    p.process(frame(p.device.layout, 0))
    p.process(frame(p.device.layout, 1, stream="new"))
    assert p.finished_reason == "device_changed"
    q = processor()
    q.process(frame(q.device.layout, 20))
    q.process(frame(q.device.layout, 0))
    assert q.finished_reason == "counter_reset"


def test_arm_rechecks_freshness_and_issues_gate(qtbot):
    session = WalkingSession(WalkingConfig())
    session.arm()
    assert session.processor is None
    layout = DeviceLayout.linear(8)
    for i in range(1000):
        session.on_frame(frame(layout, i))
    session.on_issue(AcquisitionIssue("checksum_or_tail_error", -1))
    session.arm()
    assert session.processor is None
    for i in range(1000, 2000):
        session.on_frame(frame(layout, i))
    session.preflight.context = replace(session.preflight.context, checked_monotonic_ns=1)
    session.arm()
    assert session.processor is None
    for i in range(2000, 3000):
        session.on_frame(frame(layout, i))
    session.arm()
    assert session.processor.device.layout.bit_count == 768
    assert session.processor.origin is None
    session.on_issue(AcquisitionIssue("incomplete_frame", 3001))
    assert not session.done
    session.on_issue(AcquisitionIssue("layout_mismatch", 3002))
    assert session.done


def test_configuration_and_saved_report_round_trip():
    cfg = config_from_dict({"test_type": "Sprint and Gait Test", "stop_type": "Software command", "starting_foot": "Right"})
    assert isinstance(cfg, WalkingConfig)
    assert validate_runtime_config(cfg) == []
    assert validate_runtime_config(replace(cfg, exit_clear_ms=0))
    assert validate_runtime_config(replace(cfg, stop_type="End of Time"))
    p = processor(stop_type="Software command")
    walk(p, CONTACTS, 2400)
    from data.subject_store import _report_detail, _report_from_detail
    report = p.build_report("manual")
    restored = _report_from_detail(_report_detail(report))
    assert restored.walking_summary == report.walking_summary
    assert restored.report_config_snapshot == report.report_config_snapshot


def test_walking_ui_exposes_stop_and_foot_and_dynamic_geometry(qtbot):
    from ui.param_panel import ParamPanel
    from ui.footprint_channel import FootprintChannelWidget
    from ui.views.execution_view import ExecutionView
    panel = ParamPanel()
    qtbot.addWidget(panel)
    panel.set_config(WalkingConfig(stop_type="Software command"))
    assert panel.get_config().stop_type == "Software command"
    combo = panel._widgets["stop_type"]
    assert {combo.itemText(i) for i in range(combo.count())} == {"Status change", "Software command"}
    view = ExecutionView()
    qtbot.addWidget(view)
    view.configure(WalkingConfig())
    view.on_device_state("connected", "connected")
    assert not view.btn_start.isEnabled()
    view.on_walking_readiness({"ready": True, "segment_count": 8, "message": "自检通过"})
    assert view.btn_start.isEnabled()
    channel = FootprintChannelWidget()
    qtbot.addWidget(channel)
    positions = DeviceLayout.linear(8).positions_m
    channel.render_state({"contact_bits": [0] * 767 + [1], "positions_m": positions})
    assert len(channel._contact_bits) == 768
    assert channel._y_for_index(767, 0, 100) == pytest.approx(100)


def test_parser_reports_corruption_instead_of_silent_healthy_stream():
    from hardware.protocol import ProtocolParser, crc8_poly_07
    payload = (0).to_bytes(4, "little") + bytes([1, 0]) + bytes([255]) * 12
    body = bytes([0x82]) + len(payload).to_bytes(2, "little") + payload
    wire = b'\x5a' * 4 + body + bytes([crc8_poly_07(body)]) + b'\xa5' * 4
    bad = bytearray(wire)
    bad[-5] ^= 1
    errors = []
    parser = ProtocolParser()
    parser.on_error = errors.append
    assert parser.parse(bad + wire)[0] == 0
    assert errors == ["checksum_or_tail_error"]


@pytest.mark.parametrize("running", [False, True])
def test_session_controller_requires_preflight_and_preserves_stream(qtbot, running):
    from qtpy.QtCore import QObject, Signal, Slot
    from ui.session_controller import SessionController

    class Worker(QObject):
        raw_contact_signal = Signal(list, float)
        sensor_frame_received = Signal(object)
        acquisition_issue = Signal(object)
        data_received = Signal(str)
        device_state_changed = Signal(str, str)

        def __init__(self, **kwargs):
            super().__init__()
            self.capture_starts = 0

        @Slot()
        def connect_device(self):
            self.device_state_changed.emit("connected", "connected")

        @Slot()
        def prepare_walking_capture(self):
            self.capture_starts += 1

        @Slot()
        def start_capture(self):
            pytest.fail("walking must not restart capture when armed")

        def stop(self):
            pass

    controller = SessionController(worker_factory=Worker)
    reports = []
    controller.session_finished.connect(reports.append)
    try:
        from config.overground_running_config import OvergroundRunningConfig
        config_type = OvergroundRunningConfig if running else WalkingConfig
        controller.prepare(config_type(stop_type="Software command"))
        qtbot.waitUntil(lambda: controller.device_state == "connected")
        controller.start()
        assert not controller.is_running
        worker = controller._worker
        layout = DeviceLayout.linear(3)
        for i in range(1000):
            worker.sensor_frame_received.emit(frame(layout, i))
        qtbot.waitUntil(lambda: controller._walking_ready)
        controller.start()
        qtbot.waitUntil(lambda: controller.is_running)
        assert worker.capture_starts == 1
        assert controller.engine.overground.processor.device.stream_id == "test"
        # Actual sensor frames through Qt queues, not a direct processor call.
        for i in range(1000, 3400):
            ranges = [(x - .1, x + .1) for start, end, x in CONTACTS if start <= i - 1000 < end]
            worker.sensor_frame_received.emit(frame(layout, i, ranges))
        qtbot.waitUntil(lambda: controller.engine.overground.processor.last_sample == 3399)
        worker.acquisition_issue.emit(AcquisitionIssue("layout_mismatch", 3400))
        qtbot.waitUntil(lambda: len(reports) == 1)
        assert reports[0].finish_reason == "layout_mismatch"
        summary = reports[0].running_summary if running else reports[0].walking_summary
        assert summary["valid_steps"] == 3
        assert reports[0].report_config_snapshot["device"]["segment_count"] == 3
        assert not controller.is_running
    finally:
        if controller.is_running:
            controller.stop()
        else:
            controller.discard()


def test_walking_report_shows_missing_metrics_and_restores_history(qtbot):
    from ui.views.report_view import ReportView
    from data.subject_store import _report_detail, _report_from_detail
    import json
    p = processor(1)
    walk(p, [(100, 400, .3)], 600)
    saved = json.loads(json.dumps(_report_detail(p.build_report("manual"))))
    report = _report_from_detail(saved)
    view = ReportView()
    qtbot.addWidget(view)
    view.load_report(report)
    assert "地面走路" in view._title.text()
    assert report.walking_summary["mean_stride_m"] is None
    assert len(view._replay_panel._timeline[0]["contact_bits"]) == 96


@pytest.mark.parametrize("duration,expected", [(59, 0), (60, 1)])
def test_minimum_contact_duration_includes_first_sample(duration, expected):
    p = processor(stop_type="Software command")
    walk(p, [(100, 100 + duration, .35)], 300)
    assert sum(c.confirmed for c in p.contacts) == expected
    assert (p.origin is not None) == bool(expected)


def test_export_contains_all_modules_and_walking_quality(qtbot, monkeypatch, tmp_path):
    import json
    from ui.views import report_view
    from openpyxl import load_workbook
    p = processor(8, stop_type="Software command")
    walk(p, [(100, 200, .35), (209, 800, .35)] + CONTACTS[1:], 2400)
    raw = bytes([0]) * 767 + bytes([1])
    report = p.build_report("manual", (raw,), (0.0,))
    view = report_view.ReportView()
    qtbot.addWidget(view)
    view.load_report(report)
    monkeypatch.setattr(report_view, "_get_base_dir", lambda: str(tmp_path))
    monkeypatch.setattr(report_view.QMessageBox, "question", lambda *a: report_view.QMessageBox.Yes)
    monkeypatch.setattr(report_view.QMessageBox, "information", lambda *a: None)
    errors = []
    monkeypatch.setattr(report_view.QMessageBox, "warning", lambda *a: errors.append(a))
    view._on_export()
    assert not errors
    book = load_workbook(next((tmp_path / "data").glob("*.xlsx")))
    packed = book["LED Frames"]["B2"].value.split()
    assert len(packed) == 96
    assert packed[-1] == "80"
    assert "Walking contacts" in book.sheetnames
    assert "Device and Config" in book.sheetnames
    sheet = book["Walking contacts"]
    headers = [c.value for c in sheet[1]]
    corrections = json.loads(sheet.cell(2, headers.index("merged_interruptions") + 1).value)
    assert corrections == [{"start_sample": 200, "end_sample": 209, "reason": "release_debounce"}]
    assert sheet.cell(2, headers.index("observed_samples") + 1).value == 691


def test_health_capture_handoff_does_not_reset_assembler(qtbot, monkeypatch):
    import hardware.usb_worker as usb
    from tests.test_usb_worker_lifecycle import _FakeDevice
    from hardware.protocol import UploadDataSubPack
    monkeypatch.setattr(usb, "CyUsbInterfaceDevice", _FakeDevice)
    worker = usb.UsbWorker(dll_path="fake")
    worker.connect_device()
    worker.refresh_led_health()
    initial_stream = worker._assembler.stream_id
    worker.prepare_walking_capture()
    assert worker._capturing
    assert not worker._health_check_active
    assert not worker._health_owns_capture
    assert worker._assembler.stream_id == initial_stream
    for i in range(2):
        worker._on_frame(usb.E_DATA_REPORT, None,
                         UploadDataSubPack(frameIdx=i, packNum=1, packIdx=0, buffer=bytes([255]) * 96), None)
    assert worker.layout.bit_count == 768
    worker._poll_led_health()
    assert worker._capturing
    worker.stop()


def test_one_metre_can_report_step_without_inventing_stride():
    p = processor(1, stop_type="Software command")
    walk(p, [(100, 500, .2), (400, 800, .7)], 900)
    s = p.summary()
    assert s["valid_steps"] == 1
    assert s["mean_step_m"] == pytest.approx(.5, abs=.01)
    assert s["mean_stride_m"] is None
    assert s["single_support_s"] is None


def test_separate_ambiguity_episodes_break_continuity_separately():
    p = processor(stop_type="Software command")
    for i in range(300):
        ranges = [(.1, .7)] if i < 70 or 170 <= i < 240 else []
        p.process(frame(p.device.layout, i, ranges))
    assert p.epoch == 2
    assert [x["code"] for x in p.issues] == ["ambiguous_contacts", "ambiguous_contacts"]


def test_real_usb_worker_monitor_and_stop_run_on_worker_thread(qtbot, monkeypatch):
    import hardware.usb_worker as usb
    from hardware.protocol import UploadDataSubPack
    from tests.test_usb_worker_lifecycle import _FakeDevice
    from ui.session_controller import SessionController
    monkeypatch.setattr(usb, "CyUsbInterfaceDevice", _FakeDevice)
    controller = SessionController(worker_factory=usb.UsbWorker)
    reports = []
    controller.session_finished.connect(reports.append)
    try:
        controller.prepare(WalkingConfig(stop_type="Software command"))
        qtbot.waitUntil(lambda: controller._worker._capturing)
        worker = controller._worker
        for i in range(1000):
            worker._on_frame(usb.E_DATA_REPORT, None,
                             UploadDataSubPack(frameIdx=i, packNum=1, packIdx=0, buffer=bytes([255]) * 96), None)
        qtbot.waitUntil(lambda: controller._walking_ready)
        controller.start()
        qtbot.waitUntil(lambda: controller.is_running)
        assert controller.engine.walking.processor.device.layout.bit_count == 768
        controller.stop()
        assert len(reports) == 1
        assert reports[0].report_config_snapshot["device"]["sample_rate_hz"] == 1000
    finally:
        if controller.is_running:
            controller.stop()
        else:
            controller.discard()


@pytest.mark.parametrize('segments', [3, 8, 12])
def test_short_contact_blocks_adjacency_and_releases_manual_identity(segments):
    p = processor(segments, stop_type='Software command', starting_foot='Left')
    walk(p, [(100, 400, .3), (500, 520, .8), (700, 1000, 1.3),
             (1200, 1500, 1.9), (1700, 2000, 2.5)], 2100)
    s = p.summary()
    assert [(x['from_id'], x['to_id']) for x in s['steps']] == [(2, 3), (3, 4)]
    assert len(s['cycles']) == 1
    assert len(s['contacts']) == 5
    assert s['contacts'][1]['exclusion'] == 'short_contact'
    assert s['contacts'][0]['side'] == 'left'
    assert all(c['side'] == 'unknown' for c in s['contacts'][2:])


def test_only_short_candidates_remain_reportable_before_origin():
    p = processor(stop_type='Software command')
    walk(p, [(100, 120, .3)], 200)
    report = p.build_report('manual')
    s = report.walking_summary
    assert p.origin is None
    assert s['duration_s'] is None
    assert s['valid_steps'] == 0
    assert s['contacts'][0]['exclusion'] == 'short_contact'
    assert s['excluded_contacts'] == 1


@pytest.mark.parametrize('gap', [1800, 2300])
def test_walking_stop_prefix_survives_gap_without_extending_across_it(gap):
    p = processor(stop_type='Software command')
    walk(p, [(100, 2400, .3), (2600, 2900, .9), (3100, 3400, 1.5)], 3600, skip={gap})
    s = p.summary()
    if gap == 2300:
        assert len(s['stops']) == 1
        assert s['stops'][0]['start_s'] == 0
        assert s['stops'][0]['end_s'] == pytest.approx(2.2)
    else:
        assert not s['stops']
    assert s['valid_steps'] == 1
