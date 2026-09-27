"""Independent synthetic rectangles, with known 1000-Hz event/geometry truth."""
from dataclasses import replace
import json
import time

import pytest
from config.overground_running_config import OvergroundRunningConfig
from config.overground_running_report import OvergroundRunningReport
from config.test_config import config_from_dict
from config.config_validation import validate_runtime_config
from engine.modes.overground_running_processor import OvergroundRunningProcessor
from engine.overground_session import OvergroundSession
from hardware.sensor_frame import DeviceLayout, SensorFrame, AcquisitionIssue
from hardware.beam_quality import BeamQualityPolicy
from hardware.walking_preflight import PreparedDevice


def frame(layout, n, ranges=(), received=None):
    bits = bytes(int(any(a <= x <= b for a, b in ranges)) for x in layout.positions_m)
    return SensorFrame('running', layout, n, n, received or time.perf_counter_ns(), bits,
                       bytes([1]) * layout.bit_count, b'')


def processor(n=3, **kwargs):
    layout = DeviceLayout.linear(n)
    return OvergroundRunningProcessor(OvergroundRunningConfig(**kwargs), PreparedDevice(layout, 'running', -1, 0, 1000))


def run(p, contacts, end=1300, skip=(), batch=False):
    for n in range(end):
        if n not in skip:
            ranges = [(x - .08, x + .08) for start, stop, x in contacts if start <= n < stop]
            p.process(frame(p.device.layout, n, ranges, (n // 16 + 1) * 16000000 if batch else n + 1))
    return p.summary()


TRAJECTORY = [(100, 250, .25), (400, 550, .55), (700, 850, .85)]


@pytest.mark.parametrize('segments', [1, 3, 8, 12])
@pytest.mark.parametrize('batch', [False, True])
def test_time_geometry_and_length_independent(segments, batch):
    p = processor(segments, starting_foot='Left')
    s = run(p, TRAJECTORY, batch=batch)
    assert p.origin == .1
    assert s['duration_s'] == pytest.approx(.75)
    assert s['valid_steps'] == 2 and s['valid_cycles'] == 1
    for contact in s['contacts']:
        assert contact['contact_s']['value'] == pytest.approx(.15)
        assert contact['contact_s']['valid']
    assert s['step_lengths_m'] == pytest.approx([.3, .3], abs=.0104)
    assert s['running_speed_m_s'] == pytest.approx(1, abs=.035)
    cycle = s['cycles'][0]
    assert cycle['duration_s']['value'] == pytest.approx(.6)
    assert cycle['swing_s']['value'] == pytest.approx(.45)
    assert cycle['flight_s']['value'] == pytest.approx(.3)
    assert cycle['single_support_s']['value'] == pytest.approx(.3)
    assert cycle['double_support_s']['value'] == 0
    assert s['flight_cycle_count'] == 1
    assert not p.finished_reason  # A normal mid-field clear period is not an exit.


@pytest.mark.parametrize('reverse', [False, True])
def test_cross_segment_geometry_and_direction(reverse):
    p = processor(3)
    track = [(100, 250, .98), (400, 550, 1.6), (700, 850, 2.2)]
    if reverse:
        track = [(a, b, 3 - x) for a, b, x in track]
    s = run(p, track)
    assert s['direction'] == (-1 if reverse else 1)
    assert s['step_lengths_m'] == pytest.approx([.62, .6], abs=.0104)


def test_overlap_and_zero_flight_kept_without_walking_classification():
    s = run(processor(), [(100, 450, .3), (400, 750, .7), (700, 1000, 1.1)])
    c = s['cycles'][0]
    assert c['flight_s'] == {'value': 0, 'valid': True, 'missing_reason': None}
    assert c['double_support_s']['value'] == pytest.approx(.05)
    assert c['single_support_s']['value'] == pytest.approx(.55)
    assert s['zero_flight_cycle_count'] == 1
    assert s['groups']['all']['duration_s']['count'] == 1
    assert s['groups']['with_flight']['duration_s']['count'] == 0


def test_short_clear_is_unknown_not_zero():
    s = run(processor(), [(100, 399, .3), (400, 699, .7), (700, 1000, 1.1)])
    assert not s['cycles'][0]['flight_s']['valid']
    assert s['cycles'][0]['flight_s']['missing_reason'] == 'uncertain_short_clear'
    assert s['unknown_flight_cycle_count'] == 1


def test_short_event_is_retained_and_blocks_bridging():
    s = run(processor(), [(100, 250, .3), (350, 360, .6), (500, 650, .9), (800, 950, 1.2)])
    assert len(s['contacts']) == 4
    assert s['contacts'][1]['contact_s']['missing_reason'] == 'short_contact'
    assert not s['steps'][0]['time_s']['valid']
    assert not s['steps'][1]['time_s']['valid']
    assert s['steps'][2]['time_s']['valid']
    assert all(not c['duration_s']['valid'] for c in s['cycles'])


def test_multiple_gaps_allow_recovery_without_cycles_or_side_bridge():
    p = processor(starting_foot='Left')
    track = [(100 + 300*i, 250 + 300*i, .25 + .3*i) for i in range(8)]
    s = run(p, track, 2600, skip={180, 181, 1110})
    for cycle in s['cycles']:
        if cycle['duration_s']['valid']:
            a, b = (s['contacts'][cycle[k]] for k in ('from_id', 'to_id'))
            assert a['epoch'] == b['epoch']
    assert s['valid_cycles'] > 0
    assert all(c['side'] == 'unknown' for c in s['contacts'] if c['epoch'] > 0)
    assert s['issues']


def test_boundary_toe_does_not_delete_reliable_touch_interval():
    p = processor(1)
    s = run(p, [(100, 250, .3), (400, 550, .65), (700, 850, .95)])
    # Boundary entry itself has unknown touch, while earlier reliable timings survive.
    assert s['steps'][0]['time_s']['valid']
    assert not s['contacts'][-1]['toe_m']['valid']
    assert s['contacts'][-1]['toe_m']['missing_reason'] == 'toe_clipped'
    assert s['steps'][-1]['time_s']['missing_reason'] == 'touch_not_observed'
    # Interior touch followed by toe moving onto boundary: touch-to-touch stays measurable.
    p = processor(1)
    for n in range(1000):
        ranges = [(.22, .38)] if 100 <= n < 250 else ([ (.7, .86) if n < 450 else (.84, 1)] if 400 <= n < 550 else [])
        p.process(frame(p.device.layout, n, ranges))
    s = p.summary()
    assert s['steps'][0]['time_s']['valid']
    assert not s['steps'][0]['length_m']['valid']


def test_stop_retains_duration_but_cuts_cycles():
    s = run(processor(), [(100, 2300, .3), (2400, 2550, .7), (2700, 2850, 1.1)], 3200)
    assert s['duration_s'] == pytest.approx(2.75)
    assert s['stops'] and not s['cycles'][0]['duration_s']['valid']
    assert s['steps'][1]['time_s']['valid']


def test_turn_aborts_and_acceleration_uses_weighted_speed():
    p = processor()
    s = run(p, [(100, 250, .3), (400, 550, .8), (650, 800, 1.5)])
    expected = sum(s['step_lengths_m']) / .55
    assert s['running_speed_m_s'] == pytest.approx(expected)
    p = processor()
    run(p, [(100, 250, .3), (400, 550, .9), (700, 850, .4)])
    assert p.finished_reason == 'turn_detected'


@pytest.mark.parametrize('reverse', [False, True])
def test_exit_requires_tracked_endpoint_and_500_clear_samples(reverse):
    p = processor(1)
    for n in range(1400):
        if 100 <= n < 250:
            bounds = (.22, .38)
        elif 400 <= n < 550:
            bounds = (.57, .73)
        elif 700 <= n < 850:
            bounds = (.84, 1.0)
        else:
            bounds = None
        if bounds and reverse:
            bounds = tuple(.988 - x for x in reversed(bounds))
        p.process(frame(p.device.layout, n, [bounds] if bounds else []))
        if n == 1348:
            assert p.finished_reason is None
        if n == 1349:
            assert p.finished_reason == 'passage_complete'
    assert p.summary()['duration_s'] == pytest.approx(.75)


def test_manual_end_never_auto_finishes():
    p = processor(1, stop_type='Software command')
    run(p, [(100, 250, .3), (400, 550, .65), (700, 850, .95)], 1600)
    assert p.finished_reason is None


def test_sides_need_three_per_metric():
    p = processor(8, starting_foot='Right')
    s = run(p, [(100+300*i, 250+300*i, .3+.6*i) for i in range(8)], 2500)
    side = s['sides']['step_length_m']
    assert side['left_count'] >= 3 and side['right_count'] >= 3
    assert side['asymmetry_percent']['valid']
    p = processor(starting_foot='Left')
    s = run(p, TRAJECTORY)
    assert not s['sides']['contact_s']['asymmetry_percent']['valid']


def test_preflight_config_freeze_and_fatal_changes(qtbot):
    config = config_from_dict({'test_type': 'Overground Running Test'})
    assert isinstance(config, OvergroundRunningConfig)
    assert not validate_runtime_config(config)
    session = OvergroundSession(config, OvergroundRunningProcessor)
    session.preflight.policy = BeamQualityPolicy(observation_seconds=1)
    session.arm()
    assert session.processor is None
    layout = DeviceLayout.linear(8)
    for n in range(1000):
        session.on_frame(frame(layout, n))
    session.arm()
    assert session.processor.device.layout == layout
    with pytest.raises(Exception):
        session.processor.device.stream_id = 'changed'
    session.on_frame(frame(DeviceLayout.linear(3), 1000))
    assert session.done and session.processor.finished_reason == 'device_changed'


def test_report_history_export_and_ui(qtbot):
    from data.subject_store import _report_detail, _report_from_detail, _report_summary
    from reporting.overground_running import export_sheets
    from openpyxl import Workbook
    from ui.param_panel import ParamPanel
    from ui.views.report_view import ReportView
    from ui.views.execution_view import ExecutionView
    p = processor()
    run(p, TRAJECTORY)
    report = p.build_report('manual')
    restored = _report_from_detail(json.loads(json.dumps(_report_detail(report))))
    assert isinstance(restored, OvergroundRunningReport)
    assert restored.running_summary == report.running_summary
    assert _report_summary(report)['report_type'] == 'overground_running'
    wb = Workbook()
    export_sheets(wb, restored)
    assert 'Running steps' in wb.sheetnames
    assert 'length_m.valid' in [c.value for c in wb['Running steps'][1]]
    panel = ParamPanel()
    qtbot.addWidget(panel)
    panel.set_config(OvergroundRunningConfig())
    assert panel.get_config().min_contact_time == 20
    view = ExecutionView()
    qtbot.addWidget(view)
    view.configure(OvergroundRunningConfig())
    view.on_device_state('connected', 'connected')
    assert not view.btn_start.isEnabled()
    report_view = ReportView()
    qtbot.addWidget(report_view)
    report_view.load_report(restored)
    assert '地面跑步' in report_view._title.text()


def test_no_stable_toe_keeps_times_and_best_forward_platform_is_selected():
    p = processor()
    # First stance continually changes its leading edge by > one local pitch.
    for n in range(900):
        ranges = []
        if 100 <= n < 250:
            ranges = [(.20, .30 + (.025 if n % 2 else 0))]
        elif 400 <= n < 550:
            # A short forward spike is insufficient; choose the 10-ms forward platform.
            high = .65 if 430 <= n < 440 else (.7 if n == 450 else .6)
            ranges = [(.5, high)]
        p.process(frame(p.device.layout, n, ranges))
    s = p.summary()
    assert s['contacts'][0]['toe_m']['missing_reason'] == 'no_stable_toe_platform'
    assert s['contacts'][1]['toe_m']['value'] == pytest.approx(.65, abs=.0104)
    assert s['steps'][0]['time_s']['valid']
    assert not s['steps'][0]['length_m']['valid']


def test_merged_contacts_and_noise_cannot_make_flight_or_bridge_identity():
    p = processor(starting_foot='Left')
    for n in range(1800):
        if 100 <= n < 250:
            bounds = [( .2, .35)]
        elif 400 <= n < 420:
            bounds = [(.5, .65), (.8, .95)]
        elif 420 <= n < 450:
            bounds = [(.5, .95)]  # Both tracked feet merge.
        elif 700 <= n < 850:
            bounds = [(1.1, 1.25)]
        elif 1000 <= n < 1150:
            bounds = [(1.6, 1.75)]
        elif 1300 <= n < 1450:
            bounds = [(2.1, 2.25)]
        else:
            bounds = []
        p.process(frame(p.device.layout, n, bounds))
    s = p.summary()
    assert any(i['code'] == 'ambiguous_contacts' for i in s['issues'])
    assert all(c['side'] == 'unknown' for c in s['contacts'][1:])
    assert any(c['duration_s']['valid'] for c in s['cycles'])
    assert all(c['epoch'] > 0 for c in s['contacts'] if c['id'] >= 3)


@pytest.mark.parametrize('duration,valid', [(2, False), (3, False), (19, False), (20, True)])
def test_contact_minimum_and_backdated_origin(duration, valid):
    p = processor()
    s = run(p, [(100, 100 + duration, .3)], 250)
    assert s['contacts'][0]['contact_s']['valid'] == valid
    assert p.origin == (.1 if valid else None)


def test_start_rechecks_staleness_obstruction_and_issue(qtbot):
    session = OvergroundSession(OvergroundRunningConfig(), OvergroundRunningProcessor)
    session.preflight.policy = BeamQualityPolicy(observation_seconds=1)
    layout = DeviceLayout.linear(1)
    for n in range(1000):
        session.on_frame(frame(layout, n, received=time.perf_counter_ns() - 600_000_000))
    session.arm()
    assert session.processor is None
    for n in range(1000, 2000):
        session.on_frame(frame(layout, n))
    session.on_issue(AcquisitionIssue('checksum_or_tail_error', 2000))
    session.arm()
    assert session.processor is None
    for n in range(2000, 3000):
        session.on_frame(frame(layout, n))
    session.on_frame(frame(layout, 3000, [(.2, .4)]))
    session.arm()
    assert session.processor is None


@pytest.mark.parametrize('issue', ['disconnected', 'data_timeout', 'layout_mismatch', 'out_of_order_or_reset'])
def test_fatal_session_conditions_abort(qtbot, issue):
    session = OvergroundSession(OvergroundRunningConfig(), OvergroundRunningProcessor)
    session.preflight.policy = BeamQualityPolicy(observation_seconds=1)
    for n in range(1000):
        session.on_frame(frame(DeviceLayout.linear(1), n))
    session.arm()
    if issue == 'disconnected':
        session.on_device_state('disconnected', '断连')
    elif issue == 'data_timeout':
        session.last_received_ns = time.perf_counter_ns() - 1_100_000_000
        session.poll()
    else:
        session.on_issue(AcquisitionIssue(issue, 1000))
    assert session.done
    assert session.build_report(issue).finish_reason == issue


def test_full_width_export_and_truncation_notice(qtbot, monkeypatch, tmp_path):
    from ui.views import report_view
    from openpyxl import load_workbook
    from reporting.overground_running import detail_html
    p = processor(12)
    run(p, TRAJECTORY)
    report = p.build_report('manual', (bytes(1151) + b'\x01',), (0.0,))
    report.report_config_snapshot['raw_buffer'] = {'truncated': True}
    view = report_view.ReportView()
    qtbot.addWidget(view)
    view.load_report(report)
    monkeypatch.setattr(report_view, '_get_base_dir', lambda: str(tmp_path))
    monkeypatch.setattr(report_view.QMessageBox, 'question', lambda *a: report_view.QMessageBox.Yes)
    monkeypatch.setattr(report_view.QMessageBox, 'information', lambda *a: None)
    errors = []
    monkeypatch.setattr(report_view.QMessageBox, 'warning', lambda *a: errors.append(a))
    view._on_export()
    assert not errors
    book = load_workbook(next((tmp_path / 'data').glob('*.xlsx')))
    packed = book['LED Frames']['B2'].value.split()
    assert len(packed) == 144 and packed[-1] == '80'
    assert 'Running cycles' in book.sheetnames
    assert '已截断' in detail_html(report)


def test_later_stance_gap_preserves_preceding_touch_time():
    p = processor()
    s = run(p, TRAJECTORY, skip={480})
    assert s['contacts'][1]['touch_s']['valid']
    assert not s['contacts'][1]['lift_s']['valid']
    assert not s['contacts'][1]['toe_m']['valid']
    assert s['steps'][0]['time_s']['valid']
    assert s['steps'][0]['time_s']['value'] == pytest.approx(.3)
    assert not s['steps'][1]['time_s']['valid']


def test_sqlite_history_roundtrip(tmp_path):
    from data.subject_store import SubjectStore, _report_from_detail
    from ui.views.history_view import _session_summary
    store = SubjectStore(tmp_path / 'running.sqlite3')
    p = processor()
    run(p, TRAJECTORY)
    report = p.build_report('passage_complete')
    session_id = store.record_session(None, p.config, report)
    session = store.get_session(session_id)
    restored = _report_from_detail(session.report_detail)
    assert restored.finish_reason == 'passage_complete'
    assert session.finish_reason == 'passage_complete'
    assert restored.running_summary == report.running_summary
    assert '地面跑步' in _session_summary(session)


def test_exit_evidence_cannot_survive_returning_interior():
    p = processor(1)
    for n in range(1600):
        bounds = []
        if 100 <= n < 250:
            bounds = [(.2, .35)]
        elif 400 <= n < 550:
            bounds = [(.55, .7)]
        elif 700 <= n < 850:
            bounds = [(.82, 1)] if n < 800 else [(.78, .93)]
        p.process(frame(p.device.layout, n, bounds))
    assert p.finished_reason is None


def test_cycle_support_does_not_require_future_release_outside_cycle():
    p = processor()
    # B release becomes unknown after the completed A->A interval; support up to A is known.
    s = run(p, [(100, 450, .3), (400, 950, .7), (700, 850, 1.1)], skip={900})
    assert s['cycles'][0]['flight_s']['valid']
    assert s['cycles'][0]['flight_s']['value'] == 0


def test_overlapping_inferred_same_foot_loses_identity():
    p = processor(starting_foot='Left')
    s = run(p, [(100, 850, .3), (400, 550, .7), (700, 1000, 1.1)])
    assert not s['cycles'][0]['duration_s']['valid']
    assert s['contacts'][-1]['side'] == 'unknown'


@pytest.mark.parametrize('segments', [3, 8, 12])
@pytest.mark.parametrize('reverse', [False, True])
def test_flight_requires_actual_last_support_release(segments, reverse):
    p = processor(segments)
    track = [(100, 500, .03), (300, 400, .6), (800, 950, 1.2)]
    if reverse:
        edge = max(p.positions)
        track = [(a, b, edge - x) for a, b, x in track]
    s = run(p, track)
    assert not s['contacts'][0]['lift_s']['valid']
    assert s['contacts'][1]['lift_s']['valid']
    assert s['steps'][1]['time_s']['valid']
    assert not s['steps'][1]['flight_s']['valid']
    assert s['steps'][1]['flight_s']['missing_reason'] == 'lift_not_observed'


def test_uncertain_dropout_clears_current_and_future_side_anchor():
    p = processor(8, starting_foot='Left')
    track = [(100, 180, .3), (181, 250, .3)] + [
        (400 + 300*i, 550 + 300*i, .9 + .6*i) for i in range(7)]
    s = run(p, track, 2700)
    assert s['contacts'][0]['contact_s']['missing_reason'] == 'uncertain_short_clear'
    assert all(c['side'] == 'unknown' for c in s['contacts'])
    assert s['valid_cycles'] > 0  # Subsequent A/B measurements still recover.
    assert not s['sides']['contact_s']['asymmetry_percent']['valid']


@pytest.mark.parametrize('boundary', [False, True])
def test_cycle_checks_support_from_contact_before_its_start(boundary):
    p = processor(3)
    first_position = .03 if boundary else .25
    s = run(p, [(100, 500, first_position), (300, 400, .7),
                (800, 950, 1.3), (1100, 1250, 1.9)], 1400)
    step = s['steps'][1]
    cycle = s['cycles'][1]
    assert step['time_s']['valid'] and cycle['duration_s']['valid']
    assert step['flight_s']['valid'] == (not boundary)
    assert cycle['flight_s']['valid'] == (not boundary)
    if not boundary:
        assert step['flight_s']['value'] == pytest.approx(.3)
        assert cycle['flight_s']['value'] == pytest.approx(.45)


@pytest.mark.parametrize('segments', [3, 8, 12])
def test_later_dropout_preserves_prior_step_and_cycle_times(segments):
    p = processor(segments)
    s = run(p, [(100, 250, .3), (400, 550, .7),
                (700, 780, 1.1), (781, 850, 1.1), (1000, 1150, 1.5)])
    assert s['contacts'][2]['touch_s']['valid']
    assert s['steps'][1]['time_s']['valid']
    assert s['steps'][1]['time_s']['value'] == pytest.approx(.3)
    assert s['steps'][1]['flight_s']['valid']
    assert s['cycles'][0]['duration_s']['valid']
    assert s['cycles'][0]['duration_s']['value'] == pytest.approx(.6)
    assert s['cycles'][0]['flight_s']['valid']
    assert not s['contacts'][2]['contact_s']['valid']
    assert not s['contacts'][2]['toe_m']['valid']
    assert not s['steps'][2]['time_s']['valid']
    assert not s['cycles'][1]['duration_s']['valid']


@pytest.mark.parametrize('segments', [1, 3, 8, 12])
@pytest.mark.parametrize('reverse', [False, True])
def test_entry_unknown_touch_does_not_discard_observed_flight(segments, reverse):
    p = processor(segments)
    for n in range(900):
        bounds = [(0, .16)] if 100 <= n < 150 else ([(.03, .19)] if 150 <= n < 250
                 else ([(.42, .58)] if 400 <= n < 550 else ([(.72, .88)] if 700 <= n < 850 else [])))
        if reverse:
            bounds = [(max(p.positions) - b, max(p.positions) - a) for a, b in bounds]
        p.process(frame(p.device.layout, n, bounds))
    s = p.summary()
    assert not s['contacts'][0]['touch_s']['valid']
    assert s['contacts'][0]['lift_s']['valid']
    assert not s['steps'][0]['time_s']['valid']
    assert s['steps'][0]['flight_s']['valid']
    assert s['steps'][0]['flight_s']['value'] == pytest.approx(.15)


def test_dropout_before_minimum_contact_does_not_validate_touch():
    p = processor()
    s = run(p, [(100, 250, .3), (400, 405, .7), (406, 550, .7)])
    assert not s['contacts'][1]['touch_s']['valid']
    assert not s['steps'][0]['time_s']['valid']


@pytest.mark.parametrize('gap', [1800, 2300])
def test_stop_prefix_survives_gap_without_extending_across_it(gap):
    p = processor()
    s = run(p, [(100, 2400, .3), (2600, 2750, .9), (2900, 3050, 1.5)], 3200, skip={gap})
    if gap == 2300:
        assert len(s['stops']) == 1
        assert s['stops'][0]['start_s'] == 0
        assert s['stops'][0]['end_s'] == pytest.approx(2.2)
    else:
        assert not s['stops']
    assert s['steps'][-1]['time_s']['valid']


def test_preserved_stops_do_not_duplicate_after_multiple_issues():
    p = processor()
    run(p, [(100, 2400, .3)], 2300)
    for code in ('checksum_or_tail_error', 'incomplete_frame', 'disconnected'):
        p.break_continuity(code)
    report = p.build_report('disconnected')
    assert len(report.running_summary['stops']) == 1
    assert report.running_summary['stops'][0]['end_s'] == pytest.approx(2.2)
