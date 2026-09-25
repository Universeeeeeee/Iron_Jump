from dataclasses import replace

from qtpy.QtCore import QUrl
from openpyxl import load_workbook

from config.treadmill_report import TreadmillGaitReport, TreadmillStepResult
from reporting.gait_insights import answer_gait_question
from tests.test_treadmill_processor import _summary_cycle
from ui.views.report_view import ReportView


def report():
    return TreadmillGaitReport(
        finish_reason="manual", touch_count=16, lift_count=16,
        resolved_starting_foot="left", starting_foot_source="manual",
        gait_cycles=tuple(_summary_cycle(i, "left" if i % 2 == 0 else "right", 1.0) for i in range(16)),
        per_step_results=tuple(TreadmillStepResult(
            index=i, side="left" if i % 2 == 0 else "right", row_status="valid",
            is_event_valid=True, is_included_in_statistics=True, correction_source="none",
            time_s=i + 0.6, contact_time_s=0.6, cadence_steps_per_s=2.0,
        ) for i in range(16)),
    )


def test_stable_symmetric_report_has_units_counts_and_real_references():
    r = report()
    stable = answer_gait_question(r, "步频稳定吗")
    assert "120.000 步/分钟" in stable.text
    assert "CV 0.00%" in stable.text
    assert "有效样本 16 条" in stable.text
    assert ("step", 15) in stable.references
    symmetry = answer_gait_question(r, "左右脚差异怎么样")
    assert "不对称率 0.00%" in symmetry.text
    assert "左 1.000 秒（8 条）" in symmetry.text
    assert all(kind == "cycle" for kind, _ in symmetry.references)


def test_changes_use_common_time_split_not_metric_dependent_sample_halves():
    r = report()
    r = replace(r, per_step_results=tuple(
        replace(step, cadence_steps_per_s=(2 if step.index < 8 else 3))
        for step in r.per_step_results
    ))
    answer = answer_gait_question(r, "哪里变化最大")
    assert "中点 8.000 秒" in answer.text
    assert "后减前 +60.000 步/分钟" in answer.text
    assert "步频的相对均值变化幅度最大：50.00%" in answer.text
    assert "不对称率：前半 0.00%，后半 0.00%" in answer.text


def test_asymmetric_cycles_and_insufficient_side_samples():
    r = report()
    r = replace(r, gait_cycles=tuple(
        replace(c, stance_phase_percent=70 if c.side == "right" else 60)
        for c in r.gait_cycles
    ))
    assert "不对称率 15.38%" in answer_gait_question(r, "左右差异").text
    r = replace(r, gait_cycles=r.gait_cycles[:4])
    assert "样本不足" in answer_gait_question(r, "左右差异").text


def test_cadence_does_not_require_complete_cycles():
    r = replace(report(), gait_cycles=())
    answer = answer_gait_question(r, "步频稳定吗")
    assert "120.000 步/分钟" in answer.text
    assert "步态周期有效样本 0 条" in answer.text


def test_missing_excluded_zero_and_pause_are_not_interchangeable():
    r = report()
    steps = list(r.per_step_results)
    steps[0] = replace(steps[0], cadence_steps_per_s=None)
    steps[1] = replace(steps[1], cadence_steps_per_s=float("nan"))
    steps[2] = replace(steps[2], is_included_in_statistics=False, statistics_exclusion_reason="暂停")
    steps[3] = replace(steps[3], cadence_steps_per_s=0)
    r = replace(r, per_step_results=tuple(steps), report_config_snapshot={"pause_boundaries_s": [4.3]})
    answer = answer_gait_question(r, "步频稳定吗")
    assert ("step", 0) not in answer.references
    assert ("step", 1) not in answer.references
    assert ("step", 2) not in answer.references
    assert ("step", 3) in answer.references
    assert ("step", 4) not in answer.references
    quality = answer_gait_question(r, "哪些数据没算进去")
    assert "逐步记录 3：暂停" in quality.text
    assert "2 条缺失或非有限值" in quality.text
    assert ("step", 0) in quality.references
    assert answer_gait_question(r, "具体证据呢", "stability").topic == "stability"


def test_report_question_locates_evidence_and_export_labels_simulation(qtbot, tmp_path, monkeypatch):
    import ui.views.report_view as module

    view = ReportView()
    qtbot.addWidget(view)
    r = replace(report(), report_config_snapshot={"data_source": "simulation"})
    view.load_report(r)
    answer = view.answer_gait_question("左右脚差异怎么样")
    assert "模拟演示数据" in answer
    view._locate_gait_evidence(QUrl("cycle:4"))
    assert view._cycle_detail_table.currentRow() == 4
    assert view._detail_tabs.currentWidget() is view._cycle_page
    monkeypatch.setattr(module, "_get_base_dir", lambda: str(tmp_path))
    monkeypatch.setattr(module.QMessageBox, "question", lambda *a, **k: module.QMessageBox.Yes)
    monkeypatch.setattr(module.QMessageBox, "information", lambda *a, **k: None)
    errors = []
    monkeypatch.setattr(module.QMessageBox, "warning", lambda *a, **k: errors.append(a))
    view._on_export()
    assert not errors
    files = list((tmp_path / "data").glob("*.xlsx"))
    assert len(files) == 1
    with files[0].open("rb") as stream:
        book = load_workbook(stream)
        assert "非受试者实测" in book["数据来源"].cell(1, 2).value
        assert any("Cycle" in name or "周期" in name for name in book.sheetnames)
