"""Raw acquisition, independent beam display, and dynamic monitor records."""
from dataclasses import replace
import csv
import time

import pytest
from hardware.sensor_frame import DeviceLayout, SensorFrame
from hardware.protocol import UploadDataSubPack, E_DATA_REPORT
from hardware.led_monitor import BeamDisplayFilter, MonitorWorker
from led_con_package_8m_0918.hardware.capture_diagnostics import CaptureDiagnostics
from led_con_package_8m_0918.ui.capture_archive import CaptureArchive


def frame(index, payload, *, stream='test'):
    layout = DeviceLayout.linear(len(payload) // 12)
    return SensorFrame(stream, layout, index, index, time.perf_counter_ns(),
                       layout.contact_bits(payload), bytes([1]) * layout.bit_count, payload)


def test_tail_flicker_does_not_freeze_stable_first_beam():
    display = BeamDisplayFilter(10)
    display.feed(frame(0, b'\xff' * 96))
    for i in range(1, 32):
        payload = b'\xfe' + b'\xff' * 83 + (bytes(12) if i % 2 else b'\xff' * 12)
        result = display.feed(frame(i, payload))
        assert result[0] == (0xff if i < 11 else 0xfe)
        assert result[-12:] == b'\xff' * 12
    assert sum(display.transients) > 0


def test_device_clock_and_gap_reset_confirmation():
    display = BeamDisplayFilter(10)
    display.feed(frame(0, b'\xff' * 12))
    for i in range(1, 11):
        assert display.feed(frame(i, bytes(12))) == b'\xff' * 12
    assert display.feed(frame(11, bytes(12))) == bytes(12)
    # A discontinuity cannot complete a pre-gap candidate from only two samples.
    display.feed(frame(12, b'\xff' * 12))
    assert display.feed(frame(50, bytes(12))) == bytes(12)
    assert not display.pending
    assert display.feed(frame(0, b'\xff' * 24, stream='new')) == b'\xff' * 24


@pytest.mark.parametrize('segments', [1, 3, 6, 8, 12])
def test_monitor_preserves_every_raw_frame_and_device_timestamp(qapp, tmp_path, segments):
    archive = CaptureArchive(bit_count=None)
    diagnostics = CaptureDiagnostics(tmp_path, frame_bit_count=None,
                                     frame_batch_sink=archive.append_batch)
    worker = MonitorWorker(diagnostics=diagnostics)
    unwanted = []
    worker.led_bits_signal.connect(unwanted.append)
    worker.led_contact_signal.connect(unwanted.append)
    worker.raw_contact_signal.connect(lambda *args: unwanted.append(args))
    diagnostics.start()
    payloads = [b'\xff' * (12 * segments), bytes(12 * segments), b'\xff' * (12 * segments)]
    for i, payload in enumerate(payloads):
        worker._on_frame(E_DATA_REPORT, None, UploadDataSubPack(
            bufferLen=len(payload), frameIdx=i, packNum=1, packIdx=0, buffer=payload), None)
    diagnostics.stop()
    records = list(archive.records())
    assert [ts for ts, _ in records] == [0, .001, .002]
    assert [bits for _, bits in records] == [bytes([1]) * (96 * segments), bytes(96 * segments), bytes([1]) * (96 * segments)]
    assert unwanted == []
    latest = worker.mailbox.take()
    assert len(latest.nodes) == segments
    assert latest.frames == 3
    assert latest.zero_frames == (1,) * segments
    assert diagnostics.snapshot()['dropped_frame_records'] == 0
    archive.close()


def test_long_zero_selection_keeps_original_ids_and_freezes(qapp):
    worker = MonitorWorker(observation_samples=12)
    payload = b'\xff' * 12 + bytes(12) + b'\xff' * 12
    for i in range(12):
        worker._publish_frame(frame(i, payload))
    latest = worker.mailbox.take()
    assert latest.segment_ids == ('1', '2', '3')
    assert latest.active_indices == (0, 2)
    assert latest.selection_complete
    worker._publish_frame(frame(12, b'\xff' * 36))
    assert worker.mailbox.take().active_indices == (0, 2)
    worker._publish_frame(frame(0, b'\xff' * 36, stream='next'))
    assert worker.mailbox.take().active_indices == (0, 1, 2)


def test_short_zero_does_not_change_segment_count(qapp):
    worker = MonitorWorker(observation_samples=12)
    for i in range(20):
        worker._publish_frame(frame(i, b'\xff' * 12 + (bytes(12) if i in (2, 3) else b'\xff' * 12)))
    assert worker.mailbox.take().active_indices == (0, 1)


def test_diagnostics_reports_sink_failure_and_drains_both_queues(tmp_path):
    def fail(batch):
        raise OSError('disk full')
    diagnostics = CaptureDiagnostics(tmp_path, frame_bit_count=None, frame_batch_sink=fail)
    diagnostics.record_raw(b'raw')
    diagnostics.record_frame(bytes(96), 0.0)
    diagnostics.start()
    diagnostics.stop()
    snapshot = diagnostics.snapshot()
    assert snapshot['frame_sink_errors'] == 1
    assert snapshot.get('accepted_frames', 0) == 0
    assert snapshot['last_error']
    assert not diagnostics.is_running
    rows = list(csv.DictReader(next(tmp_path.glob('raw*.csv')).open()))
    assert bytes.fromhex(rows[0]['raw_hex']) == b'raw'


def test_dynamic_archive_batch_keeps_order_and_rejects_layout_change():
    archive = CaptureArchive(bit_count=None)
    expected = [(i / 1000, bytes([i % 2]) * 288) for i in range(300)]
    archive.append_batch(expected)
    assert list(archive.records()) == expected
    with pytest.raises(ValueError):
        archive.append(.5, bytes(96))
    archive.clear()
    archive.append(0, bytes(96))
    assert archive.bit_count == 96
    archive.close()


def test_monitor_gap_delays_selection_and_resets_display(qapp):
    worker = MonitorWorker(observation_samples=12)
    for i in range(10):
        worker._publish_frame(frame(i, b'\xff' * 12 + bytes(12)))
    broken = replace(frame(20, b'\xff' * 12 + bytes(12)), dropped_frames_before=10,
                     quality_flags=('frame_gap',))
    worker._publish_frame(broken)
    for i in range(21, 31):
        worker._publish_frame(frame(i, b'\xff' * 12 + bytes(12)))
    assert not worker.mailbox.take().selection_complete


def test_ui_only_paints_latest_and_pause_keeps_health_current(qtbot):
    from led_con_package_8m_0918.ui.led_con_8m import SerialDataWidget
    widget = SerialDataWidget()
    qtbot.addWidget(widget)
    worker = MonitorWorker(mailbox=widget._display_mailbox, stability_ms=0)
    for i in range(1000):
        worker._publish_frame(frame(i, bytes([i % 2 * 255]) * 36))
    widget._refresh_ui()
    assert widget._latest.frames == 1000
    assert len(widget.led_panels) == 3
    assert widget.led_panels[0].led_states() == (1,) * 96
    assert widget._display_mailbox.take() is None
    widget.paused = True
    worker._publish_frame(frame(1000, bytes(36)))
    widget._refresh_ui()
    assert widget._latest.frames == 1001
    assert widget.led_panels[0].led_states() == (1,) * 96
    assert widget._latest.zero_frames == (501,) * 3
    widget.close()


@pytest.mark.parametrize('segments', [1, 3, 8])
def test_dynamic_ui_background_export_contains_raw_frames(qtbot, tmp_path, monkeypatch, segments):
    from qtpy.QtWidgets import QMessageBox
    from openpyxl import load_workbook
    from led_con_package_8m_0918.ui import led_con_8m as ui
    monkeypatch.setattr(ui, '_get_base_dir', lambda: str(tmp_path))
    monkeypatch.setattr(QMessageBox, 'question', lambda *a: QMessageBox.Yes)
    monkeypatch.setattr(QMessageBox, 'information', lambda *a: None)
    widget = ui.SerialDataWidget()
    qtbot.addWidget(widget)
    worker = MonitorWorker(mailbox=widget._display_mailbox)
    for i, value in enumerate((255, 0, 255)):
        raw = frame(i, bytes([value]) * (12 * segments))
        worker._publish_frame(raw)
        widget._archive.append(raw.sample_time_s, bytes([int(bool(value))]) * (96 * segments))
    widget._refresh_ui()
    assert len(widget.led_panels) == segments
    widget._save_export()
    assert not widget.btn_start.isEnabled()
    qtbot.waitUntil(lambda: widget._export_future is None)
    path = next((tmp_path / 'data').glob('*.xlsx'))
    wb = load_workbook(path, read_only=True)
    rows = list(wb.worksheets[0].values)
    assert len(rows) == 4
    assert len(rows[0]) == 96 * segments + 1
    assert [row[0] for row in rows[1:]] == [0, .001, .002]
    assert rows[2][1:] == (0,) * (96 * segments)
    assert 'Capture Info' in wb.sheetnames
    wb.close()
    widget.close()


def test_stop_waits_for_writer_without_blocking_ui(qtbot, tmp_path, monkeypatch):
    import threading
    from qtpy.QtWidgets import QMessageBox
    from led_con_package_8m_0918.ui import led_con_8m as ui
    monkeypatch.setattr(ui, '_get_base_dir', lambda: str(tmp_path))
    monkeypatch.setattr(QMessageBox, 'question', lambda *a: QMessageBox.No)
    entered, release = threading.Event(), threading.Event()
    def fake_start(worker):
        worker.diagnostics.start()
        worker._publish_frame(frame(0, b'\xff' * 12))
    monkeypatch.setattr(MonitorWorker, 'start', fake_start)
    widget = ui.SerialDataWidget()
    qtbot.addWidget(widget)
    original = widget._archive.append_batch
    def slow_append(batch):
        entered.set()
        assert release.wait(3)
        original(batch)
    monkeypatch.setattr(widget._archive, 'append_batch', slow_append)
    widget.on_start_clicked()
    try:
        qtbot.waitUntil(entered.is_set)
        widget.on_stop_clicked()
        qtbot.waitUntil(lambda: widget.serial_thread is None)
        assert widget._draining
        assert not widget.btn_start.isEnabled()
        assert not widget.btn_export.isEnabled()
        release.set()
        qtbot.waitUntil(lambda: not widget._draining)
        assert widget._archive.count == 1
        assert widget.btn_start.isEnabled()
    finally:
        release.set()
        widget.close()
        qtbot.waitUntil(lambda: widget.serial_thread is None and not widget._draining)


def test_rotating_raw_log_retains_all_chunks_and_overflow_is_visible(tmp_path):
    diagnostics = CaptureDiagnostics(tmp_path, max_csv_bytes=150, queue_capacity=8)
    expected = [bytes([i]) * 20 for i in range(8)]
    for data in expected:
        assert diagnostics.record_raw(data)
    assert not diagnostics.record_raw(b'overflow')
    diagnostics.start()
    assert diagnostics.stop()
    rows = []
    for path in sorted(tmp_path.glob('raw*.csv')):
        rows.extend(csv.DictReader(path.open()))
    assert [bytes.fromhex(row['raw_hex']) for row in rows] == expected
    assert diagnostics.snapshot()['dropped_raw_records'] == 1
    assert len(list(tmp_path.glob('raw*.csv'))) > 1


def test_monitor_reorders_physical_panels_and_records_without_extra_signals(qapp, tmp_path):
    layout = DeviceLayout.linear(3, wire_order=(2, 3, 1))
    worker = MonitorWorker(layout=layout, stability_ms=0)
    payload = b'\xff' * 12 + bytes(12) + b'\x01' * 12
    worker._on_frame(E_DATA_REPORT, None, UploadDataSubPack(36, 0, 1, 0, payload), None)
    latest = worker.mailbox.take()
    assert latest.nodes[0] == bytes([1, 0, 0, 0, 0, 0, 0, 0]) * 12
    assert latest.nodes[1] == bytes([1]) * 96
    assert latest.nodes[2] == bytes(96)
    assert latest.zero_frames == (0, 0, 1)


def test_close_during_capture_waits_and_closes_archive_after_writer(qtbot, tmp_path, monkeypatch):
    from led_con_package_8m_0918.ui import led_con_8m as ui
    monkeypatch.setattr(ui, '_get_base_dir', lambda: str(tmp_path))
    def fake_start(worker):
        worker.diagnostics.start()
        worker._publish_frame(frame(0, b'\xff' * 96))
    monkeypatch.setattr(MonitorWorker, 'start', fake_start)
    widget = ui.SerialDataWidget()
    qtbot.addWidget(widget)
    widget.show()
    widget.on_start_clicked()
    qtbot.waitUntil(lambda: widget._latest is not None)
    widget.close()
    qtbot.waitUntil(lambda: not widget.isVisible())
    assert widget.serial_thread is None
    assert not widget._diagnostics.is_running
    assert widget._archive._file.closed


def test_failed_start_can_be_retried_and_does_not_leave_writer(qtbot, tmp_path, monkeypatch):
    from led_con_package_8m_0918.ui import led_con_8m as ui
    monkeypatch.setattr(ui, '_get_base_dir', lambda: str(tmp_path))
    def fail_connect(worker):
        worker._emit_state('error', 'test device unavailable')
    monkeypatch.setattr(MonitorWorker, 'connect_device', fail_connect)
    widget = ui.SerialDataWidget()
    qtbot.addWidget(widget)
    for _ in range(10):
        widget.on_start_clicked()
        qtbot.waitUntil(lambda: widget.serial_thread is None and not widget._draining)
        assert not widget._diagnostics.is_running
        assert widget.btn_start.isEnabled()
    widget.close()


def test_diagnostics_are_not_rendered_as_ui_text(qtbot, tmp_path):
    from qtpy.QtWidgets import QLabel
    from led_con_package_8m_0918.ui.led_con_8m import SerialDataWidget
    widget = SerialDataWidget()
    qtbot.addWidget(widget)
    detail = 'Traceback: DLL C:\\internal\\driver.dll failed ACK=12 bytes=2048'
    widget._on_device_state('error', detail)
    widget._on_device_state('disconnected', 'device closed')
    widget._diagnostics = CaptureDiagnostics(tmp_path)
    widget._diagnostics._last_error = detail
    worker = MonitorWorker(mailbox=widget._display_mailbox)
    worker._publish_frame(frame(0, b'\xab' * 12))
    widget._refresh_ui()
    widget._refresh_status(time.perf_counter())
    visible = '\n'.join(label.text() for label in widget.findChildren(QLabel))
    assert '设备连接或采集失败' in visible
    for internal in ('Traceback', 'DLL', 'driver.dll', 'ACK=', 'ab ab ab'):
        assert internal not in visible
    widget.close()


def test_real_start_and_shutdown_execute_in_worker_thread(qtbot, tmp_path, monkeypatch):
    from qtpy.QtCore import QThread
    import hardware.usb_worker as shared
    from led_con_package_8m_0918.ui import led_con_8m as ui
    from tests.test_usb_worker_lifecycle import _FakeDevice
    monkeypatch.setattr(ui, '_get_base_dir', lambda: str(tmp_path))
    monkeypatch.setenv('DAYU_CAPTURE_COMMAND', '0')
    monkeypatch.setattr(shared, 'CyUsbInterfaceDevice', _FakeDevice)
    original = shared.UsbWorker.stop
    threads = []
    def stop(worker):
        threads.append(QThread.currentThread() == worker.thread())
        original(worker)
    monkeypatch.setattr(shared.UsbWorker, 'stop', stop)
    widget = ui.SerialDataWidget()
    qtbot.addWidget(widget)
    widget.on_start_clicked()
    qtbot.waitUntil(lambda: widget.serial_worker is not None and widget.serial_worker._capturing)
    assert widget.serial_worker._flush_timer.thread() == widget.serial_worker.thread()
    widget._request_stop()
    qtbot.waitUntil(lambda: widget.serial_thread is None and not widget._draining)
    assert threads == [True]
    widget.close()
